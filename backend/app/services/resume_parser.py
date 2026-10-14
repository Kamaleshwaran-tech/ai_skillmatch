"""Resume parser service using PyMuPDF (fitz) for PDF and python-docx for DOCX files.

Includes detection and explanation for empty, corrupt, or scanned (image-only) PDFs.
"""

import os
from typing import Dict, Any, List
from fastapi import HTTPException, status
import pymupdf  # PyMuPDF
import docx


class ResumeParserService:
    """Service to parse, validate, and extract textual content from uploaded resumes."""

    MIN_TEXT_LENGTH_THRESHOLD = 50  # Below this, text is likely scanned/image-only or empty

    @classmethod
    def parse_file(cls, file_path: str, filename: str) -> Dict[str, Any]:
        """Parse resume file from disk based on extension.
        
        Returns dictionary with:
        - raw_text: complete extracted text
        - page_count: number of pages (or 1 for docx)
        - sections: list of text broken by common resume headings
        - metadata: additional document details
        """
        if not os.path.exists(file_path):
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Resume file not found at: {file_path}",
            )

        file_size = os.path.getsize(file_path)
        if file_size == 0:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="The uploaded file is completely empty (0 bytes).",
            )

        ext = os.path.splitext(filename)[1].lower()

        if ext == ".pdf":
            return cls._parse_pdf(file_path)
        elif ext in [".docx", ".doc"]:
            return cls._parse_docx(file_path)
        else:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Unsupported file format '{ext}'. Only PDF (.pdf) and Word (.docx) documents are supported.",
            )

    @classmethod
    def _parse_pdf(cls, file_path: str) -> Dict[str, Any]:
        """Extract text from PDF using PyMuPDF with scanned PDF detection."""
        try:
            doc = pymupdf.open(file_path)
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Unable to read PDF file. The document may be corrupted or password-protected. Details: {str(e)}",
            )

        try:
            page_count = len(doc)
            if page_count == 0:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="The uploaded PDF contains no pages.",
                )

            extracted_pages_text = []
            has_images = False

            for page_idx in range(page_count):
                page = doc[page_idx]
                text = page.get_text("text").strip()
                if text:
                    extracted_pages_text.append(text)
                image_list = page.get_images(full=True)
                if image_list:
                    has_images = True

            combined_text = "\n\n".join(extracted_pages_text).strip()

            # Check if document has no extractable text or is a scanned image
            if len(combined_text) < cls.MIN_TEXT_LENGTH_THRESHOLD:
                if has_images:
                    raise HTTPException(
                        status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                        detail=(
                            "The PDF appears to be a scanned document or image-only file. "
                            "No selectable or extractable digital text was found. "
                            "Please upload a text-based PDF or DOCX resume created with a word processor."
                        ),
                    )
                else:
                    raise HTTPException(
                        status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                        detail=(
                            "The PDF contains insufficient readable text (under minimum threshold). "
                            "Please ensure the file is not empty or corrupted."
                        ),
                    )

            sections = cls._segment_sections(combined_text)

            return {
                "raw_text": combined_text,
                "page_count": page_count,
                "sections": sections,
                "metadata": {
                    "format": "pdf",
                    "page_count": page_count,
                    "has_images": has_images,
                    "char_count": len(combined_text),
                },
            }
        finally:
            doc.close()

    @classmethod
    def _parse_docx(cls, file_path: str) -> Dict[str, Any]:
        """Extract text from DOCX file using python-docx."""
        try:
            doc = docx.Document(file_path)
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Unable to read DOCX file. The document may be corrupted or invalid. Details: {str(e)}",
            )

        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

        # Also extract text from any tables in the docx
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    if cell_text and cell_text not in paragraphs:
                        paragraphs.append(cell_text)

        combined_text = "\n".join(paragraphs).strip()

        if len(combined_text) < cls.MIN_TEXT_LENGTH_THRESHOLD:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=(
                    "The DOCX document contains insufficient readable text. "
                    "Please ensure the document contains resume content before uploading."
                ),
            )

        sections = cls._segment_sections(combined_text)

        return {
            "raw_text": combined_text,
            "page_count": 1,
            "sections": sections,
            "metadata": {
                "format": "docx",
                "char_count": len(combined_text),
                "paragraph_count": len(paragraphs),
            },
        }

    @staticmethod
    def _segment_sections(text: str) -> List[Dict[str, Any]]:
        """Identify standard resume sections (Experience, Skills, Education, Projects)."""
        headings = [
            "skills",
            "technical skills",
            "core competencies",
            "technologies",
            "experience",
            "work experience",
            "professional experience",
            "employment history",
            "education",
            "projects",
            "certifications",
            "summary",
            "objective",
        ]

        lines = text.splitlines()
        sections = []
        current_section = "General"
        current_lines = []

        for line in lines:
            stripped = line.strip().lower().rstrip(":")
            if stripped in headings:
                if current_lines:
                    sections.append({
                        "name": current_section,
                        "text": "\n".join(current_lines).strip(),
                    })
                    current_lines = []
                current_section = line.strip()
            else:
                current_lines.append(line)

        if current_lines:
            sections.append({
                "name": current_section,
                "text": "\n".join(current_lines).strip(),
            })

        return sections
