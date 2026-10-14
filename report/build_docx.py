"""DOCX renderer for the SkillSync AI mini-project report (python-docx).

Formatting follows the conventional university specification: A4, Times New
Roman 12 pt, 1.5 line spacing, justified body text, ~1.25 inch binding margin,
lower-roman folios for the front matter and decimal folios restarting at 1 for
the body.  The table of contents, list of figures and list of tables are static
lists whose page numbers are read from pagenums.json (produced by the PDF
build), so both deliverables carry identical numbering.
"""

import json
import os

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from PIL import Image as PILImage

import ir

HERE = os.path.dirname(os.path.abspath(__file__))
FIG_DIR = os.path.join(HERE, "figures")
SHOT_DIR = os.path.join(HERE, "screenshots")
OUT = os.path.join(HERE, "output", "SkillSync_AI_Mini_Project_Report.docx")
PAGENUMS = os.path.join(HERE, "pagenums.json")

FIG_ALIASES = {"dfd0": "dfd_level0", "dfd1": "dfd_level1", "er": "er_diagram"}

TEXT_W = Cm(15.6)
GREY = RGBColor(0x54, 0x6E, 0x7A)


# ---------------------------------------------------------------------------
# low-level helpers
# ---------------------------------------------------------------------------
def _set_font(style, name="Times New Roman", size=12, bold=None, italic=None):
    style.font.name = name
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def _shade(cell, hexcolor):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor)
    tcpr.append(shd)


def _field(paragraph, instr):
    r = paragraph.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    r._r.append(fld)
    r2 = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r2._r.append(it)
    r3 = paragraph.add_run()
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    r3._r.append(fld2)
    return paragraph


def _pgnum_type(section, fmt=None, start=None):
    sectpr = section._sectPr
    pg = sectpr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        sectpr.append(pg)
    if fmt:
        pg.set(qn("w:fmt"), fmt)
    if start is not None:
        pg.set(qn("w:start"), str(start))


def _footer_page(section, size=11):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _field(p, " PAGE ")
    for r in p.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(size)
    return p


def _borderless(table):
    tbl = table._tbl
    tblpr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        el.set(qn("w:sz"), "0")
        borders.append(el)
    tblpr.append(borders)


def _placeholder_frame(doc, lines, height_cm=6.5):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    cell = t.cell(0, 0)
    cell.width = TEXT_W
    n_blanks = max(1, int(height_cm / 0.6))
    for i in range(n_blanks):
        p = cell.paragraphs[0] if i == 0 and not cell.paragraphs[0].runs else cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if i == n_blanks // 2:
            for ln in lines:
                rp = p.add_run(ln)
                rp.font.name = "Times New Roman"
                rp.font.size = Pt(10)
                rp.font.italic = True
                rp.font.color.rgb = GREY
                p.add_run().add_break()
    return t


# ---------------------------------------------------------------------------
# document scaffolding
# ---------------------------------------------------------------------------
def _setup_styles(doc):
    st = doc.styles
    _set_font(st["Normal"], size=12)
    normal = st["Normal"]
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    def mk(name, base, size, bold=False, italic=False, align=None,
           before=0, after=6, color=None):
        s = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        s.base_style = st[base]
        _set_font(s, size=size, bold=bold, italic=italic)
        pf = s.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        if align is not None:
            pf.alignment = align
        if color is not None:
            s.font.color.rgb = color
        return s

    mk("ChapTitle", "Normal", 16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
       before=6, after=18)
    mk("FrontHead", "Normal", 15, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
       before=0, after=16)
    mk("H2x", "Normal", 13.5, bold=True, before=14, after=7)
    mk("H3x", "Normal", 12, bold=True, italic=True, before=10, after=6)
    mk("Bulletx", "Normal", 12)
    mk("Captionx", "Normal", 10.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
       before=6, after=12)
    mk("TabCapx", "Normal", 10.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
       before=10, after=6)
    mk("Codex", "Normal", 8.5, before=0, after=0)
    st["Codex"].font.name = "Courier New"
    _rpr = st["Codex"].element.get_or_add_rPr()
    _rf = _rpr.find(qn("w:rFonts"))
    if _rf is None:
        _rf = OxmlElement("w:rFonts")
        _rpr.append(_rf)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        _rf.set(qn(attr), "Courier New")
    st["Codex"].paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    mk("Eqx", "Normal", 12, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER,
       before=8, after=8)
    mk("Refx", "Normal", 11, after=6)
    mk("TOC0x", "Normal", 11.5, bold=True, after=2)
    mk("TOC1x", "Normal", 11.5, after=2)
    mk("TOC2x", "Normal", 11, after=2)
    mk("Sigx", "Normal", 11.5, after=2)
    mk("PHx", "Normal", 10, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER,
       color=GREY)
    mk("TitleLine", "Normal", 14, bold=True,
       align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    mk("TitleMain", "Normal", 17, bold=True,
       align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    mk("TitleSub", "Normal", 12.5, italic=True,
       align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
    mk("TitleMid", "Normal", 11.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)


def _page_setup(section, first=False):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.4)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.4)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.3)


# ---------------------------------------------------------------------------
# block renderers
# ---------------------------------------------------------------------------
def _add_table(doc, d, with_marker_num=True):
    if d.get("num") and with_marker_num:
        p = doc.add_paragraph(f"Table {d['num']}: {d['caption']}", style="TabCapx")
    headers, rows = d["headers"], d["rows"]
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        r = p.add_run(str(h))
        r.font.bold = True
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)
        _shade(hdr[i], "E8EAF6")
    for row in rows:
        cells = t.add_row().cells
        for i, c in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(str(c))
            r.font.name = "Times New Roman"
            r.font.size = Pt(10)
    widths = d.get("widths")
    if widths:
        total = TEXT_W.cm if hasattr(TEXT_W, "cm") else 15.6
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(total * w)
    doc.add_paragraph("", style="Normal").paragraph_format.space_after = Pt(2)
    return t


def _add_figure(doc, d):
    key = d["key"]
    path = None
    if d.get("kind") == "screenshot":
        for ext in ("png", "jpg", "jpeg"):
            cand = os.path.join(SHOT_DIR, f"{key}.{ext}")
            if os.path.exists(cand):
                path = cand
                break
    else:
        base = FIG_ALIASES.get(key, key)
        for cand in (os.path.join(FIG_DIR, f"fig_{base}.png"),
                     os.path.join(FIG_DIR, f"{base}.png")):
            if os.path.exists(cand):
                path = cand
                break
    if d.get("page"):
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    if path:
        with PILImage.open(path) as im:
            pw, ph = im.size
        w = 15.0
        h = w * ph / pw
        max_h = 21.5
        if h > max_h:
            h = max_h
            w = h * pw / ph
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(path, width=Cm(w))
    else:
        if d.get("kind") == "screenshot":
            lines = ["[ SCREENSHOT TO BE INSERTED ]",
                     f"Capture this screen from the running application and save it as "
                     f"screenshots/{key}.png, then rebuild the report."]
        else:
            lines = ["[ DIAGRAM ASSET MISSING ]",
                     f"Regenerate with 'python make_diagrams.py'; expected "
                     f"figures/fig_{key}.png."]
        _placeholder_frame(doc, lines)
    doc.add_paragraph(f"Figure {d['num']}: {d['caption']}", style="Captionx")


def _add_code(doc, d):
    doc.add_paragraph(f"Listing {d['num']}: {d['caption']}", style="TabCapx")
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.cell(0, 0)
    cell.width = TEXT_W
    _shade(cell, "F5F5F5")
    lines = d["text"].split("\n")
    for i, ln in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.style = doc.styles["Codex"]
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.space_after = Pt(0)
        p.add_run(ln if ln else " ")
    doc.add_paragraph("", style="Normal").paragraph_format.space_after = Pt(2)


def _sig_table(doc, left_lines, right_lines):
    t = doc.add_table(rows=1, cols=2)
    _borderless(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, col in enumerate((left_lines, right_lines)):
        cell = t.cell(0, i)
        cell.width = Cm(7.8)
        for j, ln in enumerate(col):
            p = cell.paragraphs[0] if j == 0 else cell.add_paragraph()
            p.style = doc.styles["Sigx"]
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            p.alignment = (WD_ALIGN_PARAGRAPH.LEFT if i == 0
                           else WD_ALIGN_PARAGRAPH.RIGHT)
            r = p.add_run(ln)
            if j == 0:
                r.font.bold = True
    return t


def _toc_entries(doc, entries, pages, keyfn, style_names, leader_width):
    for level, text, _ in entries:
        pg = pages.get(keyfn(level, text), "")
        p = doc.add_paragraph(style=style_names[level])
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.tab_stops.add_tab_stop(
            Cm(15.6), 2, 1)  # right align, dot leader
        p.add_run(text)
        p.add_run("\t" + str(pg))


def _list_entries(doc, entries, pages):
    for label, cap in entries:
        pg = pages.get(label, "")
        p = doc.add_paragraph(style="TOC1x")
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.paragraph_format.tab_stops.add_tab_stop(Cm(15.6), 2, 1)
        p.add_run(f"{label}: {cap}")
        p.add_run("\t" + str(pg))


# ---------------------------------------------------------------------------
def build():
    pages = {"toc": {}, "lof": {}, "lot": {}}
    if os.path.exists(PAGENUMS):
        with open(PAGENUMS) as f:
            pages = json.load(f)["pages"]

    doc = Document()
    _setup_styles(doc)
    sec1 = doc.sections[0]
    _page_setup(sec1)

    # ---------------- title page ----------------
    tp = ir.TITLE_PAGE
    logo = os.path.join(HERE, tp["logo"])
    if os.path.exists(logo):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(logo, width=Cm(3.4))
    else:
        t = doc.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = t.cell(0, 0)
        cell.width = Cm(4.2)
        for i in range(4):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if i == 1:
                r = p.add_run("[ INSTITUTE LOGO ]")
                r.font.italic = True
                r.font.size = Pt(10)
                r.font.color.rgb = GREY
    doc.add_paragraph("")
    for line in tp["institution_lines"]:
        doc.add_paragraph(line, style="TitleLine")
    doc.add_paragraph("")
    doc.add_paragraph(tp["report_label"], style="ChapTitle")
    doc.add_paragraph(tp["title"], style="TitleMain")
    doc.add_paragraph(tp["subtitle"], style="TitleSub")
    p = doc.add_paragraph(tp["purpose"], style="TitleMid")
    doc.add_paragraph("")
    doc.add_paragraph("Submitted by", style="TitleMid")
    for name, reg in tp["student"]:
        doc.add_paragraph(f"{name}    ({reg})", style="TitleLine")
    doc.add_paragraph(tp["degree_line"], style="TitleLine")
    doc.add_paragraph(tp["year_sem"], style="TitleMid")
    doc.add_paragraph("")
    guide = tp["guide"]
    hod = tp["hod"]
    _sig_table(doc, guide, hod)
    doc.add_paragraph("")
    doc.add_paragraph(tp["department"], style="ChapTitle")
    doc.add_paragraph(tp["place_date"], style="TitleMid")

    # ---------------- front matter section ----------------
    sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
    _page_setup(sec2)
    _pgnum_type(sec2, fmt="lowerRoman", start=1)
    _footer_page(sec2)

    front_blocks, body_blocks = [], []
    seen = False
    for blk in ir.BLOCKS:
        if blk[0] == "chapter":
            seen = True
        (body_blocks if seen else front_blocks).append(blk)
    cut = next(i for i, b in enumerate(front_blocks)
               if b[0] == "front_heading" and b[1] == "LIST OF ABBREVIATIONS")
    pre, post = front_blocks[:cut], front_blocks[cut:]

    def render_front(blk, first=False):
        kind = blk[0]
        if kind == "titlepage":
            return
        if kind == "front_heading":
            if not first:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            doc.add_paragraph(blk[1], style="FrontHead")
        elif kind == "p":
            doc.add_paragraph(blk[1], style="Normal")
        elif kind == "sig_two":
            doc.add_paragraph("")
            _sig_table(doc, blk[1]["left"], blk[1]["right"])
        elif kind == "sig_one":
            doc.add_paragraph("")
            t = doc.add_table(rows=1, cols=1)
            _borderless(t)
            cell = t.cell(0, 0)
            for j, ln in enumerate(blk[1]["lines"]):
                p = cell.paragraphs[0] if j == 0 else cell.add_paragraph()
                p.style = doc.styles["Sigx"]
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                p.add_run(ln)
        elif kind == "table":
            _add_table(doc, blk[1], with_marker_num=False)

    for i, blk in enumerate(pre):
        render_front(blk, first=(i == 0))

    # ---- TOC / LOF / LOT ----
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph("TABLE OF CONTENTS", style="FrontHead")
    _toc_entries(doc, ir.TOC, pages["toc"], lambda l, t: t,
                 {0: "TOC0x", 1: "TOC1x", 2: "TOC2x"}, 15.6)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph("LIST OF FIGURES", style="FrontHead")
    _list_entries(doc, ir.LOF, pages["lof"])
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph("LIST OF TABLES", style="FrontHead")
    _list_entries(doc, ir.LOT, pages["lot"])
    for blk in post:
        render_front(blk)

    # ---------------- body section ----------------
    sec3 = doc.add_section(WD_SECTION.NEW_PAGE)
    _page_setup(sec3)
    _pgnum_type(sec3, fmt="decimal", start=1)
    _footer_page(sec3)
    hdr = sec3.header
    hdr.is_linked_to_previous = False
    hp = hdr.paragraphs[0]
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(15.6), 2)
    r = hp.add_run("AI-Driven Employment and Skill Matching Portal for Youth Development"
                   "\tSkillSync AI")
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)

    for blk in body_blocks:
        kind = blk[0]
        if kind == "chapter":
            num, title = blk[1], blk[2]
            p = doc.add_paragraph(style="ChapTitle")
            p.paragraph_format.page_break_before = True
            if isinstance(num, int):
                p.add_run(f"{num}. {title}")
            else:
                p.add_run(f"APPENDIX {num}")
                p.add_run().add_break()
                p.add_run(title)
        elif kind == "h2":
            doc.add_paragraph(blk[1], style="H2x")
        elif kind == "h3":
            doc.add_paragraph(blk[1], style="H3x")
        elif kind == "p":
            doc.add_paragraph(blk[1], style="Normal")
        elif kind == "bullets":
            for b in blk[1]:
                p = doc.add_paragraph(style="Bulletx")
                p.paragraph_format.left_indent = Cm(0.9)
                p.paragraph_format.first_line_indent = Cm(-0.45)
                p.add_run("•  " + b)
        elif kind == "nums":
            for i, b in enumerate(blk[1], 1):
                p = doc.add_paragraph(style="Bulletx")
                p.paragraph_format.left_indent = Cm(0.9)
                p.paragraph_format.first_line_indent = Cm(-0.45)
                p.add_run(f"{i}.  " + b)
        elif kind == "table":
            _add_table(doc, blk[1])
        elif kind == "fig":
            _add_figure(doc, blk[1])
        elif kind == "code":
            _add_code(doc, blk[1])
        elif kind == "eq":
            p = doc.add_paragraph(style="Eqx")
            p.add_run(f"{blk[1]}     ({blk[2]})")
        elif kind == "refs":
            p = doc.add_paragraph(style="ChapTitle")
            p.paragraph_format.page_break_before = True
            p.add_run("REFERENCES")
            for i, rtext in enumerate(blk[1], 1):
                p = doc.add_paragraph(style="Refx")
                p.paragraph_format.left_indent = Cm(0.9)
                p.paragraph_format.first_line_indent = Cm(-0.9)
                p.add_run(f"[{i}]  {rtext}")

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print("DOCX written:", build())
